#!/usr/bin/env python3
"""
Production Database vs Repository Comparison Tool

Compares SQL objects exported from production database with V0 repository
to identify manual production deployments and content differences.

Features:
- Multi-pattern version detection and history extraction
- Intelligent cosmetic change filtering
- Version comparison (Prod Higher/Lower/Same)
- Priority-based classification
- Metadata integration from reference Excel
- HTML diff generation with version highlights
- Comprehensive Excel reporting
"""

import os
import re
import hashlib
import subprocess
import sys
from collections import defaultdict
from datetime import datetime
from difflib import HtmlDiff
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Import configuration
try:
    from config import *
except ImportError:
    print("Error: config.py not found. Please ensure config.py is in the same directory.")
    sys.exit(1)

# ============================================================================
# DERIVED PATHS (from config.py)
# ============================================================================

# Script directory
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(SCRIPT_DIR, OUTPUT_SUBDIR)
OUTPUT_EXCEL = os.path.join(OUTPUT_DIR, f"{EXCEL_FILENAME_PREFIX}_{datetime.now().strftime('%Y-%m-%d_%H-%M')}.xlsx")
DIFF_DIR = os.path.join(OUTPUT_DIR, HTML_DIFF_SUBDIR)
INDEX_HTML = os.path.join(OUTPUT_DIR, SUMMARY_FILENAME)

# Create output directories
os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(DIFF_DIR, exist_ok=True)

# ============================================================================
# VERSION DETECTION PATTERNS
# ============================================================================

VERSION_PATTERNS = [
    # Standard version formats
    r'/\*\s*Version[:\s]+(\d+\.\d+(?:\.\d+)?)',  # /* Version: 1.2 */
    r'/\*\s*Ver\.?[:\s]+(\d+\.\d+(?:\.\d+)?)',   # /* Ver: 1.2 */ or /* Ver. 1.2 */
    r'/\*\s*v(\d+\.\d+(?:\.\d+)?)',              # /* v1.2 */
    r'--\s*Version[:\s]+(\d+\.\d+(?:\.\d+)?)',   # -- Version: 1.2
    
    # System-specific patterns
    r'/\*\s*PVCS\s+Version[:\s]+(\d+\.\d+)',     # /* PVCS Version: 1.2 */
    r'/\*\s*GitHub\s+Version[:\s]+(\d+\.\d+)',   # /* GitHub Version: 1.2 */
    
    # Inline version in update history (like the attached image)
    r'(\d+\.\d+)\s+(?:[A-Z]{2,5}-\d+|WMS-\d+)',  # 1.0  DevOps or 1.1  WMS-24243
]

# Version history entry patterns (for update logs)
VERSION_HISTORY_PATTERNS = [
    # Format: /* 23-Nov-2023  WLChooi  1.0  DevOps Combine Script */
    r'/\*\s*(\d{2}-[A-Za-z]{3}-\d{4})\s+(\w+)\s+(\d+\.\d+)\s+(.+?)\s*\*/',
    
    # Format: /* Date  Author  Ver.  Purposes */
    #         /* 23-Nov-2023  WLChooi  1.0  DevOps Combine Script */
    r'(\d{2}-[A-Za-z]{3}-\d{4})\s+(\w+)\s+(\d+\.\d+)\s+(.+?)(?:\*/|$)',
    
    # Format: -- 23-Nov-2023  WLChooi  1.0  DevOps Combine Script
    r'--\s*(\d{2}-[A-Za-z]{3}-\d{4})\s+(\w+)\s+(\d+\.\d+)\s+(.+?)$',
]

# ============================================================================
# COSMETIC CHANGE PATTERNS
# ============================================================================

COSMETIC_PATTERNS = {
    'set_statements': [
        r'SET\s+ANSI_NULLS\s+(ON|OFF)',
        r'SET\s+QUOTED_IDENTIFIER\s+(ON|OFF)',
        r'SET\s+ANSI_PADDING\s+(ON|OFF)',
        r'SET\s+NOCOUNT\s+(ON|OFF)',
        r'SET\s+CONCAT_NULL_YIELDS_NULL\s+(ON|OFF)',
        r'SET\s+ANSI_WARNINGS\s+(ON|OFF)',
    ],
    'permissions': [
        r'GRANT\s+.*',
        r'DENY\s+.*',
        r'REVOKE\s+.*',
    ],
    'create_alter': [
        r'CREATE\s+(PROCEDURE|FUNCTION|TRIGGER|VIEW|TABLE|SEQUENCE)',
        r'ALTER\s+(PROCEDURE|FUNCTION|TRIGGER|VIEW|TABLE|SEQUENCE)',
    ],
    'database_context': [
        r'USE\s+\[.*?\]',
    ],
    'batch_terminators': [
        r'^\s*GO\s*$',
    ],
    'schema_binding': [
        r'WITH\s+SCHEMABINDING',
    ]
}

# ============================================================================
# PRIORITY AND CLASSIFICATION DEFINITIONS
# ============================================================================

PRIORITY_LEVELS = {
    "MANUAL_PROD_DEPLOYMENT_NEWER_VERSION": {
        "priority": 1,
        "risk_level": "CRITICAL",
        "review_needed": "YES",
        "color": "FF0000",
        "description": "Production has newer version with manual deployment"
    },
    "SAME_VERSION_DIFFERENT_CONTENT": {
        "priority": 1,
        "risk_level": "CRITICAL",
        "review_needed": "YES",
        "color": "FF0000",
        "description": "Same version but content differs - manual change"
    },
    "PROD_VERSION_HIGHER_INVESTIGATE": {
        "priority": 2,
        "risk_level": "HIGH",
        "review_needed": "YES",
        "color": "FF6600",
        "description": "Production version higher - investigate"
    },
    "NO_VERSION_IN_PROD": {
        "priority": 5,
        "risk_level": "MEDIUM",
        "review_needed": "YES",
        "color": "FFC107",
        "description": "No version info in production file"
    },
    "NO_VERSION_IN_REPO": {
        "priority": 5,
        "risk_level": "MEDIUM",
        "review_needed": "YES",
        "color": "FFC107",
        "description": "No version info in repository file"
    },
    "NO_VERSION_INFO": {
        "priority": 6,
        "risk_level": "MEDIUM",
        "review_needed": "YES",
        "color": "FF9800",
        "description": "No version info in either file"
    },
    "CONTENT_DIFF_NO_VERSION": {
        "priority": 7,
        "risk_level": "MEDIUM",
        "review_needed": "YES",
        "color": "FF9800",
        "description": "Content differs but no version info"
    },
    "REPO_NEWER_NOT_DEPLOYED": {
        "priority": 8,
        "risk_level": "LOW",
        "review_needed": "NO",
        "color": "95E1D3",
        "description": "Repo has newer version - not yet deployed"
    },
    "REPO_HIGHER_VERSION": {
        "priority": 9,
        "risk_level": "LOW",
        "review_needed": "NO",
        "color": "95E1D3",
        "description": "Repository version is higher"
    },
    "SAME_VERSION_COSMETIC_ONLY": {
        "priority": 10,
        "risk_level": "NONE",
        "review_needed": "LOW",
        "color": "E8F5E9",
        "description": "Same version - only cosmetic differences"
    },
    "IDENTICAL": {
        "priority": 11,
        "risk_level": "NONE",
        "review_needed": "NO",
        "color": "C8E6C9",
        "description": "Files are identical"
    },
    "ONLY_IN_PROD": {
        "priority": 3,
        "risk_level": "HIGH",
        "review_needed": "YES",
        "color": "FF3366",
        "description": "File exists only in production"
    },
    "ONLY_IN_REPO": {
        "priority": 12,
        "risk_level": "NONE",
        "review_needed": "NO",
        "color": "E3F2FD",
        "description": "File exists only in repository"
    }
}

# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def detect_encoding(filepath):
    """Detect file encoding"""
    encodings = ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']
    for encoding in encodings:
        try:
            with open(filepath, encoding=encoding) as f:
                f.read()
            return encoding
        except (UnicodeDecodeError, UnicodeError):
            continue
    return 'binary/unknown'

def read_text_normalized(filepath):
    """Read file content with normalized line endings"""
    encodings = ['utf-8-sig', 'utf-8', 'latin-1', 'cp1252']
    for encoding in encodings:
        try:
            with open(filepath, encoding=encoding) as f:
                content = f.read()
            # Normalize line endings
            content = content.replace('\r\n', '\n').replace('\r', '\n')
            return content
        except (UnicodeDecodeError, UnicodeError):
            continue
    
    # Fallback: read as binary and decode with errors='ignore'
    with open(filepath, 'rb') as f:
        content = f.read().decode('utf-8', errors='ignore')
    content = content.replace('\r\n', '\n').replace('\r', '\n')
    return content

def normalize_sql_for_comparison(sql_text, mode='logic'):
    """
    Normalize SQL text for comparison
    mode='logic': Remove all cosmetic differences
    mode='full': Keep everything
    """
    if mode == 'full':
        return sql_text
    
    normalized = sql_text
    
    # Remove SET statements
    for pattern in COSMETIC_PATTERNS['set_statements']:
        normalized = re.sub(pattern, '', normalized, flags=re.IGNORECASE)
    
    # Remove GRANT/DENY/REVOKE
    for pattern in COSMETIC_PATTERNS['permissions']:
        normalized = re.sub(pattern, '', normalized, flags=re.IGNORECASE)
    
    # Normalize CREATE/ALTER to just CREATE
    normalized = re.sub(r'ALTER\s+', 'CREATE ', normalized, flags=re.IGNORECASE)
    
    # Remove database context
    for pattern in COSMETIC_PATTERNS['database_context']:
        normalized = re.sub(pattern, '', normalized, flags=re.IGNORECASE)
    
    # Remove batch terminators
    for pattern in COSMETIC_PATTERNS['batch_terminators']:
        normalized = re.sub(pattern, '', normalized, flags=re.IGNORECASE | re.MULTILINE)
    
    # Remove SQL comments
    normalized = re.sub(r'--.*$', '', normalized, flags=re.MULTILINE)
    normalized = re.sub(r'/\*.*?\*/', '', normalized, flags=re.DOTALL)
    
    # Normalize whitespace
    lines = normalized.split('\n')
    lines = [line.rstrip() for line in lines]
    normalized = '\n'.join(line for line in lines if line.strip())
    
    # Normalize multiple spaces
    normalized = re.sub(r'  +', ' ', normalized)
    
    return normalized.strip().lower()

# ============================================================================
# VERSION EXTRACTION AND COMPARISON
# ============================================================================

def extract_version_history(sql_text):
    """
    Extract ALL version entries from update history
    Returns list of version objects with date, author, version, purpose
    """
    version_entries = []
    
    for pattern in VERSION_HISTORY_PATTERNS:
        matches = re.finditer(pattern, sql_text, re.MULTILINE | re.IGNORECASE)
        for match in matches:
            if len(match.groups()) >= 4:
                try:
                    version_entries.append({
                        'date': match.group(1),
                        'author': match.group(2),
                        'version': match.group(3),
                        'purpose': match.group(4).strip().rstrip('*/').strip(),
                        'line': sql_text[:match.start()].count('\n') + 1
                    })
                except:
                    pass
    
    # Remove duplicates based on version number
    seen_versions = set()
    unique_entries = []
    for entry in version_entries:
        if entry['version'] not in seen_versions:
            seen_versions.add(entry['version'])
            unique_entries.append(entry)
    
    return unique_entries

def get_latest_version(version_entries):
    """Get the highest version number from history"""
    if not version_entries:
        return None
    
    # Sort by version number (semantic versioning)
    try:
        sorted_versions = sorted(
            version_entries,
            key=lambda x: [int(n) for n in x['version'].split('.')],
            reverse=True
        )
        return sorted_versions[0]['version']
    except:
        return version_entries[0]['version'] if version_entries else None

def parse_version(version_str):
    """Parse version string to list of integers"""
    if not version_str:
        return None
    try:
        return [int(x) for x in version_str.split('.')]
    except:
        return None

def compare_version_numbers(prod_version, repo_version):
    """
    Compare two version numbers
    Returns: {
        'comparison': 'Prod Higher' | 'Repo Higher' | 'Same' | 'Prod Missing' | 'Repo Missing',
        'difference': version difference string,
        'prod_numeric': parsed production version,
        'repo_numeric': parsed repo version
    }
    """
    if not prod_version and not repo_version:
        return {
            'comparison': 'Both Missing',
            'difference': 'N/A',
            'prod_numeric': None,
            'repo_numeric': None
        }
    
    if not prod_version:
        return {
            'comparison': 'Prod Missing',
            'difference': 'N/A',
            'prod_numeric': None,
            'repo_numeric': parse_version(repo_version)
        }
    
    if not repo_version:
        return {
            'comparison': 'Repo Missing',
            'difference': 'N/A',
            'prod_numeric': parse_version(prod_version),
            'repo_numeric': None
        }
    
    # Parse versions
    prod_parts = parse_version(prod_version)
    repo_parts = parse_version(repo_version)
    
    if not prod_parts or not repo_parts:
        return {
            'comparison': 'Unknown',
            'difference': 'N/A',
            'prod_numeric': prod_parts,
            'repo_numeric': repo_parts
        }
    
    # Pad to same length
    max_len = max(len(prod_parts), len(repo_parts))
    prod_parts.extend([0] * (max_len - len(prod_parts)))
    repo_parts.extend([0] * (max_len - len(repo_parts)))
    
    # Compare
    for i in range(max_len):
        if prod_parts[i] > repo_parts[i]:
            diff = sum((prod_parts[j] - repo_parts[j]) * (10 ** (max_len - j - 1)) for j in range(max_len))
            return {
                'comparison': 'Prod Higher',
                'difference': f"+{diff / (10 ** (max_len - 1)):.1f}",
                'prod_numeric': prod_parts,
                'repo_numeric': repo_parts
            }
        elif prod_parts[i] < repo_parts[i]:
            diff = sum((repo_parts[j] - prod_parts[j]) * (10 ** (max_len - j - 1)) for j in range(max_len))
            return {
                'comparison': 'Repo Higher',
                'difference': f"-{diff / (10 ** (max_len - 1)):.1f}",
                'prod_numeric': prod_parts,
                'repo_numeric': repo_parts
            }
    
    return {
        'comparison': 'Same',
        'difference': '0',
        'prod_numeric': prod_parts,
        'repo_numeric': repo_parts
    }

# ============================================================================
# FILE SCANNING
# ============================================================================

def get_files_by_name(base_path):
    """Scan directory and organize files by filename"""
    files = defaultdict(set)
    
    if not os.path.exists(base_path):
        print(f"Warning: Path does not exist: {base_path}")
        return files
    
    for root, _, filenames in os.walk(base_path):
        for filename in filenames:
            if filename.endswith('.sql'):
                rel_path = os.path.relpath(os.path.join(root, filename), base_path).replace("\\", "/")
                files[filename].add(rel_path)
    
    return files

# ============================================================================
# MAIN COMPARISON LOGIC
# ============================================================================

def classify_difference(prod_sql, repo_sql, prod_path, repo_path):
    """
    Classify the difference between production and repo files
    Returns classification dictionary
    """
    # Extract version histories
    prod_history = extract_version_history(prod_sql)
    repo_history = extract_version_history(repo_sql)
    
    prod_latest = get_latest_version(prod_history)
    repo_latest = get_latest_version(repo_history)
    
    version_comp = compare_version_numbers(prod_latest, repo_latest)
    
    # Determine production-only and repo-only versions
    prod_versions = {v['version'] for v in prod_history}
    repo_versions = {v['version'] for v in repo_history}
    prod_only = sorted(prod_versions - repo_versions, key=lambda x: parse_version(x) or [0])
    repo_only = sorted(repo_versions - prod_versions, key=lambda x: parse_version(x) or [0])
    
    # Normalize for logic comparison
    prod_logic = normalize_sql_for_comparison(prod_sql, mode='logic')
    repo_logic = normalize_sql_for_comparison(repo_sql, mode='logic')
    
    content_diff = prod_sql != repo_sql
    logic_diff = prod_logic != repo_logic
    cosmetic_only = content_diff and not logic_diff
    
    # Classification based on version comparison
    if version_comp['comparison'] == 'Prod Higher':
        if prod_only:
            classification = 'MANUAL_PROD_DEPLOYMENT_NEWER_VERSION'
            notes = f"🔴 CRITICAL: Production version {prod_latest} is HIGHER than repo {repo_latest}. Production-only versions: {', '.join(prod_only)}. Likely manual deployment without repo commit."
        else:
            classification = 'PROD_VERSION_HIGHER_INVESTIGATE'
            notes = f"⚠️ Production version {prod_latest} is HIGHER than repo {repo_latest}. Investigate."
    
    elif version_comp['comparison'] == 'Repo Higher':
        if repo_only:
            classification = 'REPO_NEWER_NOT_DEPLOYED'
            notes = f"✅ Repository version {repo_latest} is HIGHER than production {prod_latest}. Repo-only versions: {', '.join(repo_only)}. Normal - awaiting deployment."
        else:
            classification = 'REPO_HIGHER_VERSION'
            notes = f"Repository version {repo_latest} is HIGHER than production {prod_latest}."
    
    elif version_comp['comparison'] == 'Same':
        if logic_diff:
            classification = 'SAME_VERSION_DIFFERENT_CONTENT'
            notes = f"🔴 CRITICAL: Both have version {prod_latest} but CONTENT DIFFERS. Manual production change without version increment!"
        elif cosmetic_only:
            classification = 'SAME_VERSION_COSMETIC_ONLY'
            notes = f"Same version {prod_latest}, only cosmetic differences (SET statements, formatting, etc.)"
        else:
            classification = 'IDENTICAL'
            notes = f"✅ Files are identical. Version: {prod_latest}"
    
    elif version_comp['comparison'] == 'Prod Missing':
        if logic_diff:
            classification = 'NO_VERSION_IN_PROD'
            notes = f"Production file has no version info (repo: {repo_latest}). Content differs."
        else:
            classification = 'NO_VERSION_IN_PROD'
            notes = f"Production file has no version info (repo: {repo_latest}). Content matches."
    
    elif version_comp['comparison'] == 'Repo Missing':
        if logic_diff:
            classification = 'NO_VERSION_IN_REPO'
            notes = f"Repository file has no version info (prod: {prod_latest}). Content differs."
        else:
            classification = 'NO_VERSION_IN_REPO'
            notes = f"Repository file has no version info (prod: {prod_latest}). Content matches."
    
    else:  # Both Missing
        if logic_diff:
            classification = 'CONTENT_DIFF_NO_VERSION'
            notes = "Neither file has version information. Content differs - manual comparison needed."
        elif cosmetic_only:
            classification = 'SAME_VERSION_COSMETIC_ONLY'
            notes = "No version info. Only cosmetic differences."
        else:
            classification = 'IDENTICAL'
            notes = "No version info. Files are identical."
    
    priority_info = PRIORITY_LEVELS.get(classification, PRIORITY_LEVELS['NO_VERSION_INFO'])
    
    return {
        'classification': classification,
        'priority': priority_info['priority'],
        'risk_level': priority_info['risk_level'],
        'review_needed': priority_info['review_needed'],
        'notes': notes,
        'prod_version': prod_latest or '',
        'repo_version': repo_latest or '',
        'version_comparison': version_comp['comparison'],
        'version_difference': version_comp['difference'],
        'prod_version_count': len(prod_history),
        'repo_version_count': len(repo_history),
        'prod_only_versions': ', '.join(prod_only) if prod_only else '',
        'repo_only_versions': ', '.join(repo_only) if repo_only else '',
        'content_diff': 'YES' if content_diff else 'NO',
        'logic_diff': 'YES' if logic_diff else 'NO',
        'cosmetic_only': 'YES' if cosmetic_only else 'NO',
        'color': priority_info['color']
    }

# ============================================================================
# MAIN EXECUTION PLACEHOLDER
# ============================================================================

if __name__ == '__main__':
    print("=" * 80)
    print("PRODUCTION DATABASE vs REPOSITORY COMPARISON TOOL")
    print("=" * 80)
    print("\nConfiguration loaded successfully!")
    print(f"V0 Repository: {V0_REPO_BASE}")
    print(f"Production Exports: {PROD_EXPORT_BASE}")
    print(f"Reference Excel: {REFERENCE_EXCEL}")
    print(f"Output Directory: {OUTPUT_DIR}")
    print("\nReady to implement comparison logic...")
    print("\nNext steps:")
    print("1. Update PROD_EXPORT_BASE path with actual production export location")
#  GIT OPERATIONS
# ============================================================================

def switch_branch_and_pull(repo_path, branch):
    """Switch to specified branch and pull latest changes"""
    try:
        print(f"\n📂 Switching to branch '{branch}' in V0 repository...")
        
        # Checkout the branch
        checkout_result = subprocess.run(
            ["git", "checkout", branch],
            cwd=repo_path,
            capture_output=True,
            text=True
        )
        
        if checkout_result.returncode == 0:
            print(f"   ✓ Successfully switched to branch '{branch}'")
        else:
            print(f"   ⚠ Could not checkout branch '{branch}': {checkout_result.stderr.strip()}")
            print(f"   Continuing with current branch...")
            return False
        
        print(f"   📥 Pulling latest changes...")
        
        # Fetch all remote changes
        subprocess.run(
            ["git", "fetch", "--all"],
            cwd=repo_path,
            capture_output=True,
            text=True,
            check=True
        )
        
        # Pull the current branch
        result = subprocess.run(
            ["git", "pull"],
            cwd=repo_path,
            capture_output=True,
            text=True
        )
        
        if result.returncode == 0:
            print(f"   ✓ {result.stdout.strip()}")
        else:
            print(f"   ⚠ Could not pull changes: {result.stderr.strip()}")
        
        return True
    except subprocess.CalledProcessError as e:
        print(f"   ❌ Error: {e.stderr if hasattr(e, 'stderr') else str(e)}")
        print(f"   Continuing with current state...")
        return False

# ============================================================================
# METADATA LOADING
# ============================================================================

def load_reference_metadata(reference_excel_path):
    """Load metadata from reference Excel file"""
    metadata = {}
    
    if not os.path.exists(reference_excel_path):
        print(f"⚠ Reference Excel not found: {reference_excel_path}")
        return metadata
    
    try:
        print(f"\n📋 Loading reference metadata from Excel...")
        wb = load_workbook(reference_excel_path, read_only=True)
        
        # Process each sheet
        for sheet_name in wb.sheetnames:
            if 'Comparison' not in sheet_name:
                continue
            
            ws = wb[sheet_name]
            
            # Get headers
            headers = [cell.value for cell in ws[1]]
            col_map = {header: idx for idx, header in enumerate(headers, 1)}
            
            # Extract metadata for each file
            for row_idx in range(2, ws.max_row + 1):
                row_cells = list(ws[row_idx])
                
                filename = row_cells[0].value
                if not filename:
                    continue
                
                metadata[filename] = {
                    'Category': row_cells[col_map.get('Category', 7) - 1].value if col_map.get('Category') else '',
                    'Domain': row_cells[col_map.get('Domain', 8) - 1].value if col_map.get('Domain') else '',
                    'Dev Domain': row_cells[col_map.get('Dev Domain', 9) - 1].value if col_map.get('Dev Domain') else '',
                    'Unified': row_cells[col_map.get('Unified', 10) - 1].value if col_map.get('Unified') else '',
                    'Unified Date': row_cells[col_map.get('Unified Date', 11) - 1].value if col_map.get('Unified Date') else '',
                }
        
        wb.close()
        print(f"   ✓ Loaded metadata for {len(metadata)} files")
        
    except Exception as e:
        print(f"   ⚠ Error loading reference Excel: {str(e)}")
    
    return metadata

# ============================================================================
# HTML DIFF GENERATION
# ============================================================================

def generate_html_diff(prod_sql, repo_sql, filename, prod_history, repo_history, classification_info):
    """Generate HTML diff with version highlighting"""
    
    prod_lines = [line.rstrip() for line in prod_sql.splitlines()]
    repo_lines = [line.rstrip() for line in repo_sql.splitlines()]
    
    # Version summary HTML
    version_summary = f"""
    <div class="version-summary">
        <h2>Version Analysis</h2>
        <table class="info-table">
            <tr>
                <th>Production Latest Version:</th>
                <td><strong>{classification_info['prod_version'] or 'Not Found'}</strong></td>
            </tr>
            <tr>
                <th>Repository Latest Version:</th>
                <td><strong>{classification_info['repo_version'] or 'Not Found'}</strong></td>
            </tr>
            <tr>
                <th>Version Comparison:</th>
                <td class="version-comp-{classification_info['version_comparison'].replace(' ', '-').lower()}">
                    <strong>{classification_info['version_comparison']}</strong>
                </td>
            </tr>
            <tr>
                <th>Version Difference:</th>
                <td>{classification_info['version_difference']}</td>
            </tr>
            <tr>
                <th>Production Version Count:</th>
                <td>{classification_info['prod_version_count']}</td>
            </tr>
            <tr>
                <th>Repository Version Count:</th>
                <td>{classification_info['repo_version_count']}</td>
            </tr>
        </table>
        
        {f'<div class="alert alert-danger"><strong>⚠️ Production-Only Versions:</strong> {classification_info["prod_only_versions"]}</div>' if classification_info['prod_only_versions'] else ''}
        {f'<div class="alert alert-info"><strong>ℹ️ Repository-Only Versions:</strong> {classification_info["repo_only_versions"]}</div>' if classification_info['repo_only_versions'] else ''}
        
        <div class="alert alert-{classification_info['risk_level'].lower()}">
            <strong>{classification_info['classification']}</strong><br>
            {classification_info['notes']}
        </div>
    </div>
    """
    
    # Generate diff
    differ = HtmlDiff(tabsize=4, wrapcolumn=100)
    diff_html = differ.make_file(
        prod_lines,
        repo_lines,
        fromdesc="Production Database",
        todesc="V0 Repository (master)",
        context=True,
        numlines=3
    )
    
    # Enhanced CSS
    enhanced_css = """
    <style>
        body { font-family: 'Segoe UI', Arial, sans-serif; margin: 20px; background: #f5f5f5; }
        .version-summary { background: white; padding: 20px; margin-bottom: 20px; border-radius: 8px; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }
        .info-table { width: 100%; border-collapse: collapse; margin: 15px 0; }
        .info-table th { text-align: left; padding: 10px; background: #f8f9fa; border: 1px solid #dee2e6; width: 250px; }
        .info-table td { padding: 10px; border: 1px solid #dee2e6; }
        .alert { padding: 15px; margin: 15px 0; border-radius: 4px; border-left: 4px solid; }
        .alert-danger { background: #f8d7da; border-color: #dc3545; color: #721c24; }
        .alert-info { background: #d1ecf1; border-color: #0dcaf0; color: #055160; }
        .alert-critical { background: #ff6b6b; color: white; font-weight: bold; }
        .alert-high { background: #ffc107; color: #333; }
        .alert-medium { background: #ff9800; color: white; }
        .alert-low { background: #95e1d3; color: #333; }
        .alert-none { background: #c8e6c9; color: #333; }
        .version-comp-prod-higher { background: #ff6b6b; color: white; font-weight: bold; padding: 5px 10px; border-radius: 4px; }
        .version-comp-repo-higher { background: #95e1d3; color: #333; font-weight: bold; padding: 5px 10px; border-radius: 4px; }
        .version-comp-same { background: #e8f5e9; color: #333; padding: 5px 10px; border-radius: 4px; }
        table.diff { font-family: 'Courier New', monospace; border-collapse: collapse; width: 100%; background: white; box-shadow: 0 2px 4px rgba(0,0,0,0.1); }
        .diff_header { background: #343a40; color: white; padding: 10px; }
        .diff_next { background: #007bff; color: white; }
        td.diff_header { text-align: right; padding: 5px 10px; }
        .diff_add { background: #d4edda; }
        .diff_chg { background: #fff3cd; }
        .diff_sub { background: #f8d7da; }
    </style>
    """
    
    # Insert version summary and enhanced CSS
    diff_html = diff_html.replace('<head>', f'<head>{enhanced_css}')
    diff_html = diff_html.replace('</head>', f'</head><body>{version_summary}')
    diff_html = diff_html.replace('<body>', '')
    
    return diff_html

# ============================================================================
# EXCEL REPORT GENERATION
# ============================================================================

def create_excel_report(comparison_results_by_type, metadata, output_path):
    """Create Excel report with separate sheets for each object type"""
    
    print(f"\n📊 Creating Excel report with separate sheets...")
    
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)
    
    # Create a sheet for each object type
    for object_type, comparison_results in comparison_results_by_type.items():
        print(f"   Creating sheet: {object_type}")
        
        ws = wb.create_sheet(title=f"{object_type} Comparison")
        
        # Headers
        ws.append(EXCEL_COLUMNS)
        
        # Format header row
        for col_num in range(1, len(EXCEL_COLUMNS) + 1):
            cell = ws.cell(row=1, column=col_num)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Add data rows
        for result in sorted(comparison_results, key=lambda x: (x['priority'], x['filename'])):
            # Get metadata
            file_metadata = metadata.get(result['filename'], {})
            
            row_data = [
                result['filename'],
                result['status'],
                result['prod_path'],
                result['repo_path'],
                result['prod_version'],
                result['repo_version'],
                result['version_comparison'],
                result['version_difference'],
                result['prod_version_count'],
                result['repo_version_count'],
                result['prod_only_versions'],
                result['repo_only_versions'],
                result['content_diff'],
                result['logic_diff'],
                result['cosmetic_only'],
                result['classification'],
                result['priority'],
                result['risk_level'],
                result['review_needed'],
                result['html_diff_path'],
                file_metadata.get('Category', ''),
                file_metadata.get('Domain', ''),
                file_metadata.get('Dev Domain', ''),
                file_metadata.get('Unified', ''),
                file_metadata.get('Unified Date', ''),
                result['notes'],
                ''  # Description - empty for manual entry
            ]
            
            row_num = ws.max_row + 1
            ws.append(row_data)
            
            # Apply color coding
            if APPLY_EXCEL_FORMATTING:
                # Color code version comparison column
                version_comp_cell = ws.cell(row=row_num, column=7)
                if result['version_comparison'] == 'Prod Higher':
                    version_comp_cell.fill = PatternFill(start_color="FF6B6B", end_color="FF6B6B", fill_type="solid")
                    version_comp_cell.font = Font(color="FFFFFF", bold=True)
                elif result['version_comparison'] == 'Repo Higher':
                    version_comp_cell.fill = PatternFill(start_color="95E1D3", end_color="95E1D3", fill_type="solid")
                elif result['version_comparison'] == 'Same':
                    version_comp_cell.fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
                
                # Color code risk level
                risk_cell = ws.cell(row=row_num, column=18)
                risk_colors = {
                    'CRITICAL': 'FF0000',
                    'HIGH': 'FF6600',
                    'MEDIUM': 'FFC107',
                    'LOW': '95E1D3',
                    'NONE': 'C8E6C9'
                }
                if result['risk_level'] in risk_colors:
                    risk_cell.fill = PatternFill(start_color=risk_colors[result['risk_level']], 
                                                 end_color=risk_colors[result['risk_level']], 
                                                 fill_type="solid")
                    if result['risk_level'] in ['CRITICAL', 'HIGH', 'MEDIUM']:
                        risk_cell.font = Font(color="FFFFFF", bold=True)
                
                # Hyperlink for HTML diff
                html_cell = ws.cell(row=row_num, column=20)
                if result['html_diff_path']:
                    html_cell.hyperlink = result['html_diff_path']
                    html_cell.value = "View Diff"
                    html_cell.style = "Hyperlink"
        
        # Auto-fit columns
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width
        
        # Freeze header row
        ws.freeze_panes = "A2"
    
    # Save workbook
    wb.save(output_path)
    print(f"   ✓ Excel report saved: {output_path}")
    
    return output_path

# ============================================================================
# MAIN COMPARISON LOGIC
# ============================================================================

def compare_object_type(object_type, metadata):
    """Compare production and repo files for a specific object type"""
    
    print(f"\n{'='*80}")
    print(f"Processing: {object_type}")
    print(f"{'='*80}")
    
    # Map plural export folder names to singular repo folder names
    repo_folder_map = {
        'Functions': 'Function',
        'Triggers': 'Trigger',
        'Sequences': 'Sequence',
        'StoredProc': 'StoredProc',
        'Views': 'Views',
        'Tables': 'Tables'
    }
    
    prod_base = os.path.join(PROD_EXPORT_BASE, object_type)
    repo_folder = repo_folder_map.get(object_type, object_type)
    repo_base = os.path.join(V0_REPO_BASE, repo_folder)
    
    print(f"Production: {prod_base}")
    print(f"Repository: {repo_base}")
    
    # Scan directories
    prod_files = get_files_by_name(prod_base)
    repo_files = get_files_by_name(repo_base)
    
    all_filenames = sorted(set(prod_files.keys()) | set(repo_files.keys()))
    
    print(f"Production files: {len(prod_files)}")
    print(f"Repository files: {len(repo_files)}")
    print(f"Total unique files: {len(all_filenames)}")
    
    results = []
    diff_count = 0
    
    # Create subdirectory for this object type's diffs
    object_diff_dir = os.path.join(DIFF_DIR, object_type)
    os.makedirs(object_diff_dir, exist_ok=True)
    
    for idx, filename in enumerate(all_filenames, 1):
        if VERBOSE and idx % 50 == 0:
            print(f"   Processing {idx}/{len(all_filenames)}...")
        
        in_prod = filename in prod_files
        in_repo = filename in repo_files
        
        result = {
            'filename': filename,
            'status': '',
            'prod_path': '',
            'repo_path': '',
            'html_diff_path': '',
            'priority': 99,
            'classification': '',
            'risk_level': '',
            'review_needed': '',
            'notes': '',
            'prod_version': '',
            'repo_version': '',
            'version_comparison': '',
            'version_difference': '',
            'prod_version_count': 0,
            'repo_version_count': 0,
            'prod_only_versions': '',
            'repo_only_versions': '',
            'content_diff': 'N/A',
            'logic_diff': 'N/A',
            'cosmetic_only': 'N/A',
            'color': 'FFFFFF'
        }
        
        if in_prod and in_repo:
            result['status'] = 'BOTH'
            prod_paths = list(prod_files[filename])
            repo_paths = list(repo_files[filename])
            result['prod_path'] = '; '.join(prod_paths)
            result['repo_path'] = '; '.join(repo_paths)
            
            # Read files
            prod_full_path = os.path.join(prod_base, prod_paths[0])
            repo_full_path = os.path.join(repo_base, repo_paths[0])
            
            prod_sql = read_text_normalized(prod_full_path)
            repo_sql = read_text_normalized(repo_full_path)
            
            # Classify difference
            classification_info = classify_difference(prod_sql, repo_sql, prod_full_path, repo_full_path)
            
            result.update(classification_info)
            
            # Generate HTML diff if content differs
            if classification_info['content_diff'] == 'YES' and ENABLE_HTML_DIFFS:
                prod_history = extract_version_history(prod_sql)
                repo_history = extract_version_history(repo_sql)
                
                html_filename = f"{filename.replace('.sql', '')}.html"
                html_path = os.path.join(object_diff_dir, html_filename)
                
                html_content = generate_html_diff(prod_sql, repo_sql, filename, 
                                                  prod_history, repo_history, classification_info)
                
                with open(html_path, 'w', encoding='utf-8') as f:
                    f.write(html_content)
                
                result['html_diff_path'] = os.path.join(object_type, html_filename)
                diff_count += 1
        
        elif in_prod:
            result['status'] = 'ONLY PROD'
            result['prod_path'] = '; '.join(list(prod_files[filename]))
            result['classification'] = 'ONLY_IN_PROD'
            priority_info = PRIORITY_LEVELS['ONLY_IN_PROD']
            result['priority'] = priority_info['priority']
            result['risk_level'] = priority_info['risk_level']
            result['review_needed'] = priority_info['review_needed']
            result['notes'] = 'File exists only in production database'
            result['color'] = priority_info['color']
        
        else:  # in_repo only
            result['status'] = 'ONLY REPO'
            result['repo_path'] = '; '.join(list(repo_files[filename]))
            result['classification'] = 'ONLY_IN_REPO'
            priority_info = PRIORITY_LEVELS['ONLY_IN_REPO']
            result['priority'] = priority_info['priority']
            result['risk_level'] = priority_info['risk_level']
            result['review_needed'] = priority_info['review_needed']
            result['notes'] = 'File exists only in repository'
            result['color'] = priority_info['color']
        
        results.append(result)
    
    print(f"   ✓ Completed: {len(results)} files processed, {diff_count} HTML diffs generated")
    
    return results

# ============================================================================
# MAIN EXECUTION
# ============================================================================

def main():
    print("=" * 80)
    print("PRODUCTION DATABASE vs REPOSITORY COMPARISON TOOL")
    print("=" * 80)
    
    # Check configuration
    if PROD_EXPORT_BASE == "/path/to/production/exports":
        print("\n❌ ERROR: Please update PROD_EXPORT_BASE in config.py")
        print("   Set it to the path where production database exports are located")
        return
    
    if not os.path.exists(PROD_EXPORT_BASE):
        print(f"\n❌ ERROR: Production export path does not exist: {PROD_EXPORT_BASE}")
        print("   Please export production database objects first using PowerShell scripts")
        return
    
    print(f"\n📋 Configuration:")
    print(f"   Production Exports: {PROD_EXPORT_BASE}")
    print(f"   V0 Repository: {V0_REPO_BASE}")
    print(f"   V0 Branch: {V0_BRANCH}")
    print(f"   Reference Excel: {REFERENCE_EXCEL}")
    print(f"   Output Directory: {OUTPUT_DIR}")
    
    # Switch to master branch and pull latest
    switch_branch_and_pull(V0_REPO_ROOT, V0_BRANCH)
    
    # Load reference metadata
    metadata = {}
    if ENABLE_METADATA_IMPORT and os.path.exists(REFERENCE_EXCEL):
        metadata = load_reference_metadata(REFERENCE_EXCEL)
    
    # Compare each object type and store results separately
    results_by_type = {}
    all_results = []
    
    for object_type in OBJECT_TYPES:
        results = compare_object_type(object_type, metadata)
        results_by_type[object_type] = results
        all_results.extend(results)
    
    # Create Excel report with separate sheets
    excel_path = create_excel_report(results_by_type, metadata, OUTPUT_EXCEL)
    
    # Generate summary statistics
    if GENERATE_SUMMARY_STATS:
        print(f"\n{'='*80}")
        print("SUMMARY STATISTICS")
        print(f"{'='*80}")
        
        total_files = len(all_results)
        critical_count = len([r for r in all_results if r['risk_level'] == 'CRITICAL'])
        high_count = len([r for r in all_results if r['risk_level'] == 'HIGH'])
        medium_count = len([r for r in all_results if r['risk_level'] == 'MEDIUM'])
        prod_higher_count = len([r for r in all_results if r['version_comparison'] == 'Prod Higher'])
        
        print(f"Total Files Compared: {total_files}")
        print(f"Critical Issues (Manual Prod Changes): {critical_count}")
        print(f"High Priority Issues: {high_count}")
        print(f"Medium Priority Issues: {medium_count}")
        print(f"Production Version Higher: {prod_higher_count}")
        print(f"\n ATTENTION: Focus on {critical_count + high_count} files requiring immediate review")
    
    print(f"\n{'='*80}")
    print("✅ COMPARISON COMPLETED")
    print(f"{'='*80}")
    print(f"\n📁 Results:")
    print(f"   Excel Report: {excel_path}")
    print(f"   HTML Diffs: {DIFF_DIR}")

if __name__ == '__main__':
    main()